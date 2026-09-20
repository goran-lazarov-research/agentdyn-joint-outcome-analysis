"""Build the English empirical manuscript and an updated Serbian protocol."""
from pathlib import Path
import json
import runpy
from copy import deepcopy
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT.parent
OUT = WORK.parent / 'itt2026_output'
helpers = runpy.run_path(str(WORK / 'analysis/build_documents.py'))
base, p, h, table = [helpers[k] for k in ['base', 'p', 'h', 'table']]
oldrefs = helpers['REFS']
R = json.loads((ROOT / 'results/extended_results.json').read_text())
T = R['totals']
REFS = [oldrefs[i] for i in [0,1,2,3,4,7,5,6,8]] + [
    ('R. B. Nelsen, An Introduction to Copulas, 2nd ed. New York, NY: Springer, 2006, doi: 10.1007/0-387-28678-0.', 'https://doi.org/10.1007/0-387-28678-0'),
    ('B. Efron, “Bootstrap Methods: Another Look at the Jackknife,” The Annals of Statistics, vol. 7, no. 1, pp. 1–26, 1979, doi: 10.1214/aos/1176344552.', 'https://doi.org/10.1214/aos/1176344552'),
    oldrefs[9], oldrefs[10]
]


def mr(text, upright=False):
    r=OxmlElement('m:r')
    if upright:
        props=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');props.append(sty);r.append(props)
    rp=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Cambria Math')
    rp.append(fonts);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');rp.append(sz);r.append(rp)
    t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=str(text);r.append(t)
    return r


def seq(value):
    if isinstance(value,str):return [mr(value)]
    if isinstance(value,list):return value
    return [value]


def sub(value,index):
    x=OxmlElement('m:sSub')
    for name,v in [('e',value),('sub',index)]:
        z=OxmlElement('m:'+name)
        for item in seq(v):z.append(item)
        x.append(z)
    return x


def sup(value,index):
    x=OxmlElement('m:sSup')
    for name,v in [('e',value),('sup',index)]:
        z=OxmlElement('m:'+name)
        for item in seq(v):z.append(item)
        x.append(z)
    return x


def frac(top,bottom):
    x=OxmlElement('m:f')
    for name,v in [('num',top),('den',bottom)]:
        z=OxmlElement('m:'+name)
        for item in seq(v):z.append(item)
        x.append(z)
    return x


def total(index,limit,body):
    n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch)
    loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'subSup');pr.append(loc);n.append(pr)
    for name,v in [('sub',index),('sup',limit),('e',body)]:
        el=OxmlElement('m:'+name)
        for item in seq(v):el.append(item)
        n.append(el)
    return n


def accent(value,char='̂'):
    a=OxmlElement('m:acc');pr=OxmlElement('m:accPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),char);pr.append(ch);a.append(pr)
    el=OxmlElement('m:e')
    for item in seq(value):el.append(item)
    a.append(el);return a


def equation(d,number,nodes):
    z=d.add_paragraph();f=z.paragraph_format
    f.first_line_indent=Inches(0);f.space_before=Pt(5);f.space_after=Pt(5)
    f.keep_together=True;f.keep_with_next=True
    f.tab_stops.add_tab_stop(Inches(1.57),WD_TAB_ALIGNMENT.CENTER)
    f.tab_stops.add_tab_stop(Inches(3.49),WD_TAB_ALIGNMENT.RIGHT)
    z.add_run('\t')
    math=OxmlElement('m:oMath')
    for node in nodes:math.append(node)
    z._p.append(math)
    run=z.add_run('\t('+str(number)+')');run.font.size=Pt(9)
    return z


def equation_rows(rows):
    arr=OxmlElement('m:eqArr')
    for nodes in rows:
        row=OxmlElement('m:e')
        for node in nodes:row.append(node)
        arr.append(row)
    return arr


def references(d,count=12,small=True):
    for i,(text,url) in enumerate(REFS[:count],1):
        z=p(d,f'[{i}] {text} {url}')
        f=z.paragraph_format;f.first_line_indent=Inches(-.16);f.left_indent=Inches(.16)
        f.space_after=Pt(3);f.keep_together=True
        z.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for run in z.runs:run.font.size=Pt(8 if small else 9)


def figure(d,name,caption):
    z=p(d,'');z.paragraph_format.first_line_indent=Inches(0)
    z.paragraph_format.keep_with_next=True
    z.add_run().add_picture(str(ROOT/'results'/name),width=Inches(3.4))
    p(d,caption,'Caption')


def manuscript():
    d=base(True)
    title=p(d,'Ranking Prompt Injection Defenses for AI Agents through Joint Outcome Bounds','Title')
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    s=d.add_section(WD_SECTION_START.CONTINUOUS)
    cols=s._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'360')
    z=p(d,'Abstract',bold=True);z.paragraph_format.first_line_indent=Inches(0)
    p(d,'Prompt injection defenses must preserve legitimate task completion while preventing attacker goals. Published evaluations often report these outcomes as separate rates, leaving their dependence unknown. We present a secondary analysis of AgentDyn tables covering 122 model–configuration combinations and three task suites, together with 288 archived execution records. Classical probability bounds identify feasible joint completion rates without assuming independence. The primary comparison includes 12 model labels, 10 configurations, and 540 within-model pairs. Using suite-level marginals instead of pooled marginals increases the number of strictly ordered pairs from 337 to 394. In 298 pairs, the ordering supported by joint completion bounds opposes the ordering based on lower attack success. An archival pilot confirms the bounds and shows that multiplying marginal rates can misestimate joint completion by 3.30 percentage points. Three inconsistencies between reported overall values and suite averages are documented; excluding the affected configurations preserves the main pattern. These are conditional comparisons of reported benchmark outcomes, not deployment guarantees. The contribution is a reproducible audit of ranking identifiability using existing data, without new model executions.')
    z=p(d,'Index Terms—AI agents, indirect prompt injection, cybersecurity, joint outcomes, partial identification, benchmark evaluation.');z.paragraph_format.first_line_indent=Inches(0);z.runs[0].italic=True

    h(d,'I Introduction')
    p(d,'An autonomous language model agent can use external information to select and execute tools. Indirect prompt injection exploits this interaction by placing adversarial instructions in content outside the original user request [1]. An evaluation must therefore distinguish completing the authorized task from satisfying the attacker’s objective. Both events can occur in one execution, and neither event may occur when the agent makes no useful progress.')
    p(d,'AgentDojo [2] and InjecAgent [3] provide controlled settings for studying such attacks. Defenses including CaMeL [4], Progent [5], and DRIFT [6] introduce controls over information flow, privileges, or execution plans. AgentDyn [7] extends evaluation to tasks requiring dynamic planning and useful environmental instructions. Its published results provide both task completion and attack success rates across multiple configurations.')
    p(d,'Those separate rates do not generally determine how often a task succeeds while its specified attack fails. Treating their product as a joint probability assumes independence. Conversely, ranking only by attack success can prefer a configuration that rarely completes tasks. The practical question is how much of a joint-outcome ranking can be established from the published evidence, even when the full paired records are unavailable.')
    p(d,'We apply established probability bounds to the published AgentDyn results and quantify the value of reporting marginals separately by task suite. The study asks three questions: which within-model comparisons can be resolved without knowing the dependence between outcomes; how often this ordering opposes the ordering by attack success alone; and how an illustrative penalty for successful attacks changes the set of possible preferred configurations. A separately identified archival pilot supplies exact joint outcomes and an exploratory uncertainty check.')

    h(d,'II Related Work and Scope')
    p(d,'Existing benchmarks and defenses already recognize that security and utility must be evaluated together [2]–[7]. The present work does not introduce that principle or a new defense. Its focus is the information needed to support comparative conclusions from the reported rates. AgentDyn itself documents utility loss in several defensive pipelines [7]; our analysis examines which joint-outcome comparisons remain logically determined when dependence information is omitted.')
    p(d,'Defense names denote full archived pipelines. Differences can include prompts, filtering, planners, model calls, or enforcement rules. A comparison between two such pipelines cannot isolate the causal effect of a particular privilege restriction. We therefore avoid attributing performance differences to policy granularity. Later revisions of a defense must also remain distinct from earlier benchmark implementations.')
    p(d,'Adaptive attackers can invalidate conclusions drawn from a fixed attack family [8]. Our results concern the evaluated objectives represented in the source tables and the pilot’s important_instructions records. They do not establish general resistance to adaptive attacks, other attack surfaces, or previously unmeasured unauthorized actions. Failure of a specified attack objective is narrower than full compliance with an authorization policy.')

    h(d,'III Data and Audit Procedure')
    p(d,'The aggregate source is AgentDyn version 3, dated May 7, 2026 [7]. We mechanically extract Tables 9–17, expanding row spans and joining rows by the exact model and defense labels. The source provides clean utility, attacked utility, and attack success for 122 configurations in Shopping, GitHub, and DailyLife. This yields 366 suite-level configuration records, each with three reported metrics. These are published aggregate observations, not new executions performed for this study.')
    p(d,'The main panel contains 12 model labels and the same 10 configurations: no defense, Prompt Sandwiching, Spotlighting, ProtectAI, PIGuard, PromptGuard2, Tool Filter, CaMeL, Progent, and DRIFT. It yields 120 conditions and 12 × 45 = 540 unordered within-model pairs. Two Meta-SecAlign rows are preserved in the data package but excluded from this rectangular comparison because their model identities do not match the 12 shared labels. The number of pairwise comparisons is not a number of independent experiments.')
    p(d,'The source describes overall performance as an equal average across three suites [7]. We therefore define the primary estimand using equal suite weights and reconstruct all overall rates from the suite cells. We retain the separately printed Overall column for an arithmetic audit, rather than silently replacing either source. A difference exceeding 0.01 percentage points between Overall and the mean of the three rounded cells is flagged because ordinary rounding to two decimals cannot explain it.')
    p(d,'The separate pilot uses 288 public logs from repository snapshot 5353cf7615b135cace8d07c8f12dac53a16b6db3 [9]. It comprises 12 user tasks, four per suite, two attacks per task, two models, and four configurations. Its 192 attacked and 96 clean records were selected with seed 20260915 before the selected outcomes were downloaded. Git blob hashes identify every source file. This is an exploratory secondary study, not an externally registered confirmatory experiment.')
    p(d,'We keep the two evidence sources separate. The aggregate tables provide broader model coverage but omit the paired outcome counts. The pilot provides paired outcomes for a smaller, earlier archived sample. We do not combine them into a single sample size or use the pilot to reconstruct unobserved joint rates for the full tables. The full raw archive was not reanalyzed; source-version and denominator alignment remain assumptions of the aggregate comparisons.')

    h(d,'IV Mathematical Model')
    h(d,'A Joint outcomes and dependence',2)
    p(d,'For a valid attacked execution i, let U be one if the authorized task completes and A be one if the specified attacker goal succeeds. The desired joint outcome is')
    equation(d,1,[sub('J','i'),mr(' = '),sub('U','i'),mr('(1 − '),sub('A','i'),mr('),   '),sub('U','i'),mr(', '),sub('A','i'),mr(' ∈ {0,1}.')])
    p(d,'Let u, a, and j be the respective mean rates over the same evaluated population and weights. The four joint cells are task success with attack failure, both events succeeding, neither event succeeding, and attack success with task failure. Their dependence gives the identity')
    equation(d,2,[mr('j = u(1 − a) − '),mr('Cov',True),mr('(U,A).')])
    p(d,'Thus the independence product is correct only when the covariance vanishes. The neither-event cell does not by itself prove protective blocking: it can include refusal, inability, or another failure mechanism. Technical errors require a separate category and must not be interpreted automatically as confirmed attacker success or confirmed protection.')
    h(d,'B Feasible joint rates',2)
    p(d,'Classical Fréchet–Hoeffding bounds [10] yield')
    equation(d,3,[mr('L = '),mr('max',True),mr('(0,u − a),   H = '),mr('min',True),mr('(u,1 − a).')])
    p(d,'The feasible joint rate lies between L and H. To see this, the probability that U and A both equal one lies between max(0, u + a − 1) and min(u, a). Subtracting that probability from u gives (3). Both endpoints are attainable by suitable joint distributions with the stated marginals. We use these established bounds as an audit tool; no new probability theorem is claimed.')
    p(d,'For each suite s, consider its lower and upper bounds and a nonnegative weight, with all weights summing to one. The joint rate of the specified mixture is bounded by')
    equation(d,4,[sub('L','w'),mr(' = '),total('s=1','3',[sub('w','s'),sub('L','s')]),mr(',   '),sub('H','w'),mr(' = '),total('s=1','3',[sub('w','s'),sub('H','s')]),mr('.')])
    p(d,'Here each weight is one third. Applying (3) only after averaging u and a produces pooled bounds. Because maximum is convex and minimum is concave, the weighted suite lower bound cannot be smaller than the pooled lower bound, and the suite upper bound cannot be larger than the pooled upper bound. Reporting suite marginals can therefore remove ambiguity without introducing an independence assumption.')
    p(d,'Each percentage is printed to two decimals. We conservatively treat it as an interval of half-width 0.005 percentage points, clipped to zero and one after conversion to proportions. Using superscripts minus and plus for its lower and upper endpoints, the rounding-aware bounds are')
    equation(d,5,[equation_rows([
        [sub('L','s'),mr(' = '),mr('max',True),mr('(0,'),sup(sub('u','s'),'−'),mr(' − '),sup(sub('a','s'),'+'),mr('),')],
        [sub('H','s'),mr(' = '),mr('min',True),mr('('),sup(sub('u','s'),'+'),mr(',1 − '),sup(sub('a','s'),'−'),mr(').')]
    ])])
    p(d,'These are identification ranges: they describe which joint rates are compatible with the marginals and rounding. They are not confidence intervals, do not estimate sampling variability, and cannot repair inconsistent populations or denominators. Both reported marginals in a suite must refer to the same evaluated cases and weighting scheme.')
    h(d,'C Pairwise ordering and decision sensitivity',2)
    p(d,'For configurations d and q within a fixed model m, we certify a strict order only when the intervals are separated:')
    equation(d,6,[mr('d ≻ q   ⇔   '),sub('L','md'),mr(' > '),sub('H','mq'),mr('.')])
    p(d,'An overlap means the ordering is unresolved by these marginals; it does not demonstrate equivalent performance. A reversal relative to attack-success ranking occurs when the lower-ASR configuration has a strictly lower joint interval. Rounding uncertainty is respected for both comparisons. We count these relations descriptively, without treating pairs sharing a model or task population as statistically independent.')
    p(d,'An illustrative decision model assigns one unit of benefit to joint completion and a penalty of lambda units to a successful attack:')
    equation(d,7,[sub('S','d'),mr('(λ) = '),sub('j','d'),mr(' − λ'),sub('a','d'),mr(',   λ ≥ 0.')])
    p(d,'Its feasible score interval is obtained by substituting the feasible joint and attack-rate endpoints in each suite and then averaging. We examine lambda values 0, 0.1, 0.25, 0.5, 1, 2, and 5. This is a sensitivity model for the evaluated attack cases, not an estimate of production attack frequency or calibrated economic loss. A configuration remains a possible winner if its score upper bound reaches the largest score lower bound.')
    h(d,'D Estimation from paired records',2)
    p(d,'Where individual records are available, the joint rate is estimated directly. Within a fixed model, denote the task count in suite s by T with subscript s. The barred J value is the mean joint outcome over the attacks on task t. The task-weighted estimator is')
    equation(d,8,[sub(accent('j'),'d'),mr(' = '),total('s=1','3',[frac(sub('w','s'),sub('T','s')),total('t=1',sub('T','s'),sub(accent('J','̄'),'dst'))]),mr('.')])
    p(d,'The pilot has four tasks per suite and two attacks per task, so this estimator equals the pooled mean over its 24 attacked records per condition. In the archived attacked logs, security=True denotes attacker-goal success; clean records follow a different convention and are excluded from this calculation. None of the selected pilot records logs a technical error.')
    p(d,'We compute paired differences between configurations on identical tasks and obtain exploratory percentile intervals from 10,000 stratified task bootstrap replicates [11], using seed 20260915. A resampled task retains both attacks and all configurations. These intervals address variation across the small empirical task sample, whereas (3)–(5) address unknown dependence. Neither procedure measures variability across fresh model executions. Degenerate all-zero or all-one bootstrap intervals are suppressed.')

    h(d,'V Results')
    h(d,'A Information gained from suite reporting',2)
    p(d,'Table I summarizes the 540 comparisons. Pooled marginals separate 337 pairs (62.4%), while suite-level marginals separate 394 (73.0%), an increase of 57 pairs or 10.6 percentage points. The remaining 146 pairs cannot be strictly ordered using the available marginals. Mean feasible interval width falls from 5.54 to 4.78 percentage points across the 120-condition panel; 51 conditions have narrower ranges.')
    z=p(d,'TABLE I  Strictly separated joint-outcome pairs. Each model has 45 configuration pairs. Reversed denotes a separated joint ordering opposed to the ordering by lower ASR. All entries are calculated from published suite cells.','Caption')
    z.paragraph_format.keep_with_next=True
    vals=[]
    for z in R['model_summaries']:
        name=z['model'].replace('Claude-Sonnet-','Claude ').replace('Gemini-2.5 ','Gemini ').replace('Qwen3 235B-A22B','Qwen3 235B')
        vals.append([name,z['pooled_resolved_pairs'],z['stratified_resolved_pairs'],z['certified_asr_reversals']])
    vals.append(['Total',337,394,298])
    table(d,['Model','Pooled','Suite','Reversed'],vals,[1.62,.58,.58,.62],small=True,keep_together=True)
    p(d,'The largest narrowing occurs for unprotected GPT-4o. Its pooled feasible joint range is 17.71%–55.53%, whereas suite information narrows it to 23.88%–48.69%, a width reduction of 13.00 percentage points. Figure 1 shows why a single exact joint score would overstate the available evidence. Several relatively productive configurations retain overlapping ranges.')
    figure(d,'joint_bounds_gpt4o.png','Fig. 1. GPT-4o joint completion ranges from pooled and suite-level published marginals. These ranges describe identification uncertainty, not 95% confidence intervals. Endpoints include rounding tolerance.')
    h(d,'B Ordering conflicts and penalty sensitivity',2)
    p(d,'Across the panel, 298 of 540 pairs (55.2%) have a certified joint ordering opposed to the ordering by lower ASR. This is an objective-dependent conflict, not evidence that preventing attacks is unimportant. Low ASR accompanied by low useful completion can be undesirable when both outcomes matter. The joint criterion also does not account for harms beyond the measured attacker goal.')
    p(d,'For GPT-4o, the archived CaMeL row reports both attacked utility and ASR as zero. The rounding-aware joint upper bound is consequently 0.005%, while the DRIFT joint lower bound is 26.25%. The latter is strictly higher despite its larger reported ASR. This comparison illustrates why a zero attack-success number alone is insufficient for selecting a useful agent configuration. It does not evaluate a current CaMeL implementation.')
    p(d,'At lambda zero, the GPT-4o marginals leave four possible joint winners: no defense, Prompt Sandwiching, Spotlighting, and DRIFT. At lambda one, two, and five, only DRIFT remains a possible winner in this particular reported panel. This result is conditional on the illustrative scoring rule and the source data. It must not be generalized to adaptive attacks, newer pipeline versions, or a production deployment with different task and loss distributions.')
    h(d,'C Exact outcomes in the archival pilot',2)
    p(d,'The pilot supplies an independent check of the accounting. All eight exact joint rates lie inside their corresponding marginal bounds. Table II shows the four joint-cell counts, making simultaneous user and attacker success visible. The independence proxy differs from the exact joint rate by as much as 3.30 percentage points; its error can have either sign.')
    p(d,'TABLE II  Pilot attacked-run counts. Columns are the values of (U,A); each row contains 24 records. Clean records are separate.','Caption')
    for model,label in [('gpt-4o-2024-08-06','GPT-4o'),('google_gemini-2.5-pro','Gemini 2.5 Pro')]:
        z=p(d,label,bold=True);z.paragraph_format.keep_with_next=True
        rows=[]
        for x in R['pilot_validation']:
            if x['model']==model:
                name={'none':'No added defense','spotlighting':'Spotlighting','tool_filter':'Tool Filter','progent':'Progent'}[x['defense']]
                rows.append([name]+[x['counts_UA'].get(key,0) for key in ['10','11','00','01']])
        table(d,['Configuration','(1,0)','(1,1)','(0,0)','(0,1)'],rows,[1.56,.46,.46,.46,.46],small=True,keep_together=True)
    p(d,'For unprotected GPT-4o, the exact joint rate is 10/24 = 41.67%. Multiplying task success 13/24 by attack failure 17/24 instead gives 38.37%. Three records satisfy both the user and attacker objectives. The joint rate of Spotlighting is 11/24 = 45.83%; its paired difference from the baseline is 4.17 percentage points, with an exploratory 95% task bootstrap interval of −8.33 to 16.67 points. For Gemini, the corresponding difference is 8.33 points with an interval of −4.17 to 20.83 points. Both intervals include zero.')
    h(d,'D Source arithmetic and sensitivity',2)
    p(d,'Three attacked-utility rows have an Overall value that differs from the equal mean of their suite cells beyond rounding tolerance. For Spotlighting with Claude-Sonnet-3.5, the two values are 58.33% and 57.09%; for Claude-Sonnet-4.5, 68.33% and 68.89%; for DRIFT with Kimi-K2.5, 22.44% and 22.55%. We preserve these differences and do not infer which source entry should be corrected.')
    p(d,'Removing all three affected configurations leaves 117 conditions and 513 within-model pairs. The separated-pair counts are then 312 from pooled marginals and 369 from suite marginals, still an increase of 57. There are 281 certified reversals relative to lower ASR. The principal observations therefore persist under this exclusion, although the comparison remains conditional on the correctness and population alignment of the other source cells.')

    h(d,'VI Discussion and Limitations')
    p(d,'The main methodological lesson is to distinguish an observable joint outcome from an assumed one. When paired records exist, report the four-cell outcome table and technical errors. When only marginals exist, bounds can preserve useful comparisons without manufacturing an exact joint rate. Reporting suite-level cells increases the information available for this purpose, even if no further model calls are possible.')
    p(d,'A benchmark report should identify the denominator, suite weights, treatment of failed evaluations, attack family, and pipeline revision for every metric. Otherwise, even correct probability algebra can be applied to incompatible rates. Our aggregate analysis assumes matched case populations within each suite because the full paired archive was not reprocessed. We cannot reconstruct its technical-error policy, within-task dependence, or repeated-run variation from the published marginals.')
    p(d,'The mathematical separation criterion is deliberately strict but does not create a statistical significance test. It conditions on reported empirical rates and asks which joint rates are compatible with them. A model’s population performance could differ on new tasks or new executions. Likewise, the 540 comparisons share configurations and source tasks and must not be interpreted as 540 independent confirmations of a hypothesis.')
    p(d,'The pilot is limited to 12 task clusters, one archived realization per condition, two models, and one attack family. The available timestamps are January 22 and January 24, 2026; package and benchmark metadata are incomplete in some records. The later Progent revision [5], for example, must not be equated with that archived implementation. Pilot intervals remain exploratory and should not be transferred to the aggregate tables.')
    p(d,'The joint objective measures successful completion with failure of one specified attack. It can miss other unauthorized actions and does not quantify the severity of a successful attack. Our lambda sensitivity analysis is an explicit choice of preference, not a universal ranking. A useful extension would publish paired outcomes under adaptive attacks and validate each tool action against an independently specified authorization policy.')

    h(d,'VII Conclusion')
    p(d,'A secondary analysis of 122 reported configurations shows how joint-outcome bounds can make defense comparisons more transparent. In the complete 120-condition panel, suite reporting resolves 57 additional within-model pairs, and 298 comparisons conflict with an ordering based only on lower ASR. Exact archived outcomes confirm that multiplying marginal rates need not recover joint completion. The practical recommendation is to publish paired outcome counts, denominators, and suite-level results, and to distinguish identification uncertainty from sampling uncertainty. These findings concern the cited archived benchmark and do not establish deployment safety.')
    h(d,'Data and Code Availability')
    p(d,'The accompanying package includes the extracted numeric source tables, original pilot manifest and 288 source logs, analysis scripts, derived bounds and comparisons, plots, source-arithmetic audit, and dependency versions. The analysis runs locally without model access. Upstream attribution and license notices are preserved. The full source article is not redistributed. No public archive identifier is claimed for this working package.')
    h(d,'Acknowledgment of AI Assistance')
    p(d,'OpenAI ChatGPT assisted with study design, mathematical exposition, drafting and editing all sections, and generating analysis and document scripts. The supplied code computed the reported results from cited public data. This disclosure follows IEEE guidance [12]; the submitting authors remain responsible for verifying the analysis, citations, and manuscript.')
    h(d,'References');references(d)
    d.add_section(WD_SECTION_START.CONTINUOUS)
    d.core_properties.title='Ranking Prompt Injection Defenses for AI Agents through Joint Outcome Bounds'
    d.save(OUT/'ITT2026_Pocetni_rukopis_EN.docx')


def protocol():
    d=base()
    p(d,'Analiza zaštite AI agenata za ITT 2026','Title')
    p(d,'Istraživački protokol i stanje rukopisa','Subtitle')
    p(d,'16. septembar 2026. | Verzija 2.0')
    p(d,'Pripremljen je engleski rukopis zasnovan na sekundarnoj analizi javno objavljenih podataka. Rad ispituje koliko se pouzdano mogu porediti zaštitne konfiguracije AI agenata kada su završavanje zadatka i uspešnost napada objavljeni kao odvojene stope. Analiza je izvršena bez novih poziva modelima, plaćenog API pristupa ili GPU resursa. Ova verzija zamenjuje raniji plan novog kontrolisanog eksperimenta.')
    h(d,'Naslov i konferencijska tema')
    p(d,'Ranking Prompt Injection Defenses for AI Agents through Joint Outcome Bounds',bold=True)
    p(d,'Tema odgovara oblasti Responsible AI, Cybersecurity and Digital Trust konferencije ITT 2026. Rok je 21. septembar 2026, a skup se održava 28. i 29. oktobra u Dubaiju. Poziv navodi šest strana za regularne radove i četiri do šest za kratke radove, IEEE šablon i dvostruko anonimnu recenziju. Najavljena IEEE tehnička podrška i indeksiranje još nisu potvrđeni.')
    h(d,'Podaci i izvedena analiza')
    p(d,'Šira analiza koristi numeričke tabele 9–17 iz treće verzije rada AgentDyn od 7. maja 2026. Izdvojene su 122 kombinacije modela i konfiguracije u tri okruženja: Shopping, GitHub i DailyLife. Za svaku su dostupne stope uspešnosti bez napada, uspešnosti pod napadom i uspešnosti napada. Glavno poređenje obuhvata 12 modela sa istih 10 konfiguracija, odnosno 120 kombinacija i 540 parova unutar istog modela. Dve preostale kombinacije sa drugačijim modelima čuvaju se u podacima, ali se ne uključuju u ovo poređenje.')
    p(d,'Odvojena pilot analiza koristi 288 ranije objavljenih zapisa: 192 izvršavanja pod napadom i 96 bez napada. Uzorak sadrži 12 različitih korisničkih zadataka, dva modela i četiri konfiguracije. To nisu 288 nezavisnih zadataka. Svaki fajl je proveren prema identitetu sadržaja u fiksiranom javnom repozitorijumu. Cela arhiva izvršavanja nije ponovo obrađena; zaključci iz širokih tabela uslovljeni su usklađenošću njihovih populacija i imenilaca.')
    table(d,['Pokazatelj','Izračunata vrednost'],[
        ['Parovi razdvojeni granicama iz ukupnih proseka','337 od 540'],
        ['Parovi razdvojeni granicama po okruženjima','394 od 540'],
        ['Dodatno razjašnjena poređenja','57 parova'],
        ['Suprotan poredak u odnosu na nižu stopu napada','298 parova'],
        ['Neslaganja ukupne vrednosti i proseka okruženja','3 reda izvornih tabela'],
        ['Nova izvršavanja modela','0']
    ],[5.2,1.7])
    d.add_page_break()
    h(d,'Matematički sadržaj rukopisa')
    p(d,'U engleskom rukopisu nalazi se osam numerisanih formula u izvornom Word formatu. Prvo se definišu binarni ishodi korisničkog zadatka U i cilja napadača A, a zatim zajednički uspeh J = U(1 − A). Formula sa kovarijansom objašnjava zašto proizvod odvojenih stopa ne mora da bude zajednička stopa uspeha.')
    p(d,'Klasične Fréchet–Hoeffding granice daju najnižu i najvišu moguću zajedničku stopu na osnovu margina. Granice se računaju posebno za svako okruženje i zatim proseče jednakim težinama. Posebna formula obuhvata zaokruživanje objavljenih procenata. Dve konfiguracije strogo se porede samo ako su odgovarajući intervali razdvojeni. Preklapanje intervala znači da dostupne margine ne određuju poredak.')
    p(d,'Model osetljivosti S(λ) = j − λa prikazuje uticaj izabrane kazne za uspešan napad. Parametar λ je analitička pretpostavka, a ne procena ekonomskog gubitka. Poslednja formula definiše procenu iz uparenih zapisa. Pilot koristi 10.000 ponovnih uzorkovanja celih zadataka po okruženjima. Ti eksplorativni intervali pouzdanosti jasno su odvojeni od granica mogućih zajedničkih stopa.')
    h(d,'Doprinos i granice zaključivanja')
    p(d,'Doprinos je ponovljiva empirijska provera koliko objavljeni podaci određuju rangiranje zaštita. Ne tvrdi se da su probabilističke granice novi matematički rezultat niti da je razvijena nova zaštita. Kada se uklone tri konfiguracije sa aritmetičkim neslaganjima, ostaje 513 parova, sa 312 razdvojenih poređenja iz ukupnih proseka i 369 iz podataka po okruženjima. Dobitak od 57 poređenja ostaje isti.')
    p(d,'Rezultati se odnose na navedene istorijske verzije sistema i evaluirane ciljeve napadača. Ne dokazuju bezbednost u produkciji, otpornost na sve adaptivne napade ili uzročni efekat pojedinačnog mehanizma zaštite. Veća citiranost ne može da se garantuje. Otvoreni podaci, jasan matematički okvir i ponovljiv kod mogu učiniti rad korisnim drugim istraživačima.')
    h(d,'Šta je potrebno pre predaje')
    p(d,'Autorska provera treba da obuhvati tumačenje izvora, matematičku notaciju, usklađenost stopa po okruženjima i zaključke. Dokument ima konferencijski raspored u dve kolone, ali konačan raspored treba preneti u zvanični šablon koji dostavi organizator. Rukopis je anonimizovan. Izjava o AI pomoći mora ostati tačna, a konačni autori odgovorni za sadržaj.')
    p(d,'Za ovu sekundarnu studiju nije potreban novi API pristup. Dodatna izvršavanja predstavljala bi buduće proširenje. Prateći paket sadrži ulaze, skripte, rezultate, verzije biblioteka i uputstvo za lokalno ponavljanje analize. Javna objava koda ili predaja konferenciji nisu izvršene.')
    h(d,'Izvori za nastavak')
    for text in [
        'Poziv ITT 2026: https://easychair.org/cfp/ITT-2026',
        'AgentDyn v3: https://arxiv.org/abs/2602.03117v3',
        'Javni zapisi: https://github.com/SaFo-Lab/AgentDyn',
        'IEEE smernice za AI pomoć: https://open.ieee.org/author-guidelines-for-artificial-intelligence-ai-generated-text/'
    ]:
        z=p(d,text);z.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for run in z.runs:run.font.size=Pt(9)
    d.core_properties.title='Analiza zaštite AI agenata za ITT 2026'
    d.save(OUT/'ITT2026_Istrazivacki_protokol.docx')


if __name__=='__main__':
    manuscript()
    protocol()
    print(OUT/'ITT2026_Pocetni_rukopis_EN.docx')
    print(OUT/'ITT2026_Istrazivacki_protokol.docx')
