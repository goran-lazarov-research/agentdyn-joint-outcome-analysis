"""Prepare the edited empirical text and figure for the final A4 template builder.

Run apply_itt2026_template.py for the deliverable. This intermediate stage retains
the earlier IEEE-style layout; the final builder imports the user-supplied A4
template, adds all authors, and replaces this stage's temporary output/report.
"""
from pathlib import Path
from copy import deepcopy
import json
import re
import runpy
import hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT.parent
OUT=WORK.parent/'itt2026_output'
SOURCE=OUT/'ITT2026_Pocetni_rukopis_EN.docx'
TARGET=OUT/'ITT2026_Final_Manuscript_EN.docx'
legacy=runpy.run_path(str(ROOT/'analysis/build_empirical_manuscript.py'))
if not SOURCE.exists():legacy['manuscript']()
d=Document(SOURCE)
R=json.loads((ROOT/'results/extended_results.json').read_text())
mr,sub,sup,total,frac,accent,eqrows=[legacy[x] for x in ['mr','sub','sup','total','frac','accent','equation_rows']]

def find(starts):
    return next(p for p in d.paragraphs if p.text.startswith(starts))

def replace(starts,text):
    p=find(starts);p.text=text;return p

def after(anchor,text,style='Normal'):
    el=OxmlElement('w:p');anchor._p.addnext(el)
    p=Paragraph(el,anchor._parent);p.style=d.styles[style];p.add_run(text)
    return p

def inline(p,parts):
    p.clear()
    for x in parts:
        if isinstance(x,str):p.add_run(x)
        else:
            m=OxmlElement('m:oMath');m.append(x);p._p.append(m)

def eq_after(anchor,nodes):
    p=after(anchor,'')
    p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.keep_together=True
    p.paragraph_format.keep_with_next=True
    p.add_run('\t')
    m=OxmlElement('m:oMath')
    for x in nodes:m.append(x)
    p._p.append(m);p.add_run('\t(0)')
    return p

def move_table_after(anchor,headers,rows,widths):
    t=legacy['table'](d,headers,rows,widths,small=True,keep_together=True)
    anchor._p.addnext(t._tbl)
    # The shared helper adds a trailing spacer at the end of the document.
    tail=d.paragraphs[-1]
    if not tail.text and not tail._p.xpath('.//w:sectPr'):
        tail._p.getparent().remove(tail._p)
    return t

title='Ranking Prompt Injection Defenses for AI Agents Using Joint Outcome Bounds'
find('Ranking Prompt').text=title
abstract=find('Prompt injection defenses must')
abstract.text=(
    'Abstract—Prompt injection defenses must preserve legitimate task completion while preventing attacker goals. '
    'Published evaluations often report these outcomes separately, leaving their dependence unknown. '
    'We analyze AgentDyn tables covering 122 model–configuration combinations in three task suites and, separately, '
    '288 archived execution records. Classical probability bounds identify feasible joint completion rates without '
    'assuming independence. The main comparison includes 12 model labels, 10 configurations, and 540 within-model '
    'pairs. Suite-level marginals increase the number of strictly ordered pairs from 337 to 394 compared with pooled '
    'marginals. In 298 pairs, the joint ordering opposes the ordering by lower attack success. The archival pilot '
    'illustrates why multiplying marginal rates can misestimate joint completion, with an observed error of up to '
    '3.30 percentage points. We document three inconsistencies between printed overall values and suite averages; '
    'excluding the affected configurations preserves the increase of 57 ordered pairs. A penalty sensitivity analysis '
    'makes preference assumptions explicit. These conditional benchmark comparisons provide a reproducible audit '
    'of ranking identifiability without new model executions. They do not establish deployment safety.')
old_abstract=find('Abstract')
if old_abstract is not abstract and old_abstract.text=='Abstract':
    old_abstract._p.getparent().remove(old_abstract._p)
replace('AgentDojo [2] and InjecAgent',
    'AgentDojo [2] and InjecAgent [3] provide controlled settings for studying such attacks. Defenses including '
    'CaMeL [4], Progent [5], and DRIFT [6] introduce controls over information flow, privileges, or execution plans. '
    'AgentDyn [7] extends evaluation to dynamic planning and useful environmental instructions. Its reported '
    'task-completion rates and attack success rates (ASRs) support a secondary analysis across archived pipelines.')
replace('Those separate rates',
    'Those separate rates do not generally determine how often a task succeeds while its specified attack fails. '
    'Multiplying task success by attack failure assumes independence. Ranking only by ASR can prefer a configuration '
    'that rarely completes tasks. The practical question is how much of a joint-outcome ranking can be established '
    'from published evidence when paired records are unavailable.')
p=after(find('We apply established probability bounds'),
    'The contributions are an explicit identification criterion for comparing reported pipelines, an empirical '
    'audit of the information gained from suite-level reporting, and a reproducibility package connecting aggregate '
    'bounds to exact archival outcomes. The method uses established probability results; its contribution lies in '
    'their application and the resulting audit, rather than a new defense or probability theorem.')
p=after(find('Adaptive attackers can invalidate'),
    'The threat model concerns adversarial instructions placed in material an agent encounters while pursuing a '
    'legitimate task. External content may also contain useful instructions, so rejecting all such content can '
    'degrade task completion. We evaluate the recorded user and attacker objectives; we do not assume that the '
    'attacker controls the trusted user request or that every failure to reach an attacker objective proves that '
    'the execution remained within an authorization policy. This distinction is especially relevant when autonomous '
    'agents act on retrieved information rather than only generate a textual answer.')
p=after(find('The separate pilot uses'),
    'The pilot selection is paired across configurations. Four task identifiers are sampled in each suite and '
    'two injection objectives are sampled for each selected task. The same identifiers are used for GPT-4o '
    '(gpt-4o-2024-08-06) and Gemini 2.5 Pro (google_gemini-2.5-pro), under no added defense, '
    'spotlighting_with_delimiting, tool_filter, and progent. Each condition therefore contains 24 attacked records '
    'and 12 clean records. Selection, timestamps, and content hashes are retained in the manifest.')
p=after(find('We keep the two evidence sources separate'),
    'The extraction checks uniqueness of every model–configuration key, equality of key sets across metrics, '
    'valid percentage ranges, and completeness of the 12-by-10 panel. The analysis then computes bounds for '
    'each suite, averages them with stated weights, compares unordered configurations within each model, and '
    'repeats the comparison after excluding arithmetic flags. An offline rerun reproduced all five numerical '
    'JSON outputs exactly with the recorded software versions. Neither extraction nor reproduction executes '
    'the instructions stored in the source logs.')

# Mathematical clarity and explicit endpoint construction.
p=replace('Let u, a, and j be',
    'Let u, a, and j be the mean rates of U, A, and J over the same evaluated population and weights. '
    'Throughout the equations, rates are proportions; tables and figures express them as percentages. '
    'The dependence between task and attack outcomes gives the identity')
p=replace('The feasible joint rate lies',
    'The feasible joint rate lies in [L,H]. To see this, the probability that U and A both equal one lies '
    'between max(0,u+a−1) and min(u,a). Subtracting this probability from u gives (3). More explicitly, '
    'let p with subscripts x and y denote the joint probability of U=x and A=y. For any feasible j, '
    'the four cells can be constructed as')
p=eq_after(p,[eqrows([
    [sub('p','10'),mr(' = j,   '),sub('p','11'),mr(' = u − j,')],
    [sub('p','01'),mr(' = a − u + j,   '),sub('p','00'),mr(' = 1 − a − j.')]
])])
p=after(p,
    'These cells sum to one and are nonnegative exactly when L ≤ j ≤ H. Thus both endpoints are attainable. '
    'The bounds cannot be tightened using only those two marginals. This is an application of classical bounds '
    '[10], not a claim that the construction is a new result.')
p=find('For each suite s, consider')
inline(p,['For suite s, let ',sub('w','s'),' be a nonnegative weight, with weights summing to one. '
          'Using the suite-specific lower and upper bounds, the mixture is bounded by'])
p=replace('Here each weight is one third.',
    'Here each weight is one third. Applying (3) after averaging u and a produces pooled bounds. Convexity '
    'of the maximum and concavity of the minimum imply that the suite lower bound cannot be smaller and its '
    'upper bound cannot be larger. Additional suite information can therefore remove ambiguity without an '
    'independence assumption. The weighted endpoints are attainable because each suite can realize its endpoint '
    'separately. They describe the specified equal-suite mixture, rather than a deployment mixture with unknown weights.')
p=replace('Its feasible score interval is obtained',
    'For the rounding-aware suite endpoints, the score interval is')
p=eq_after(p,[eqrows([
    [sup('S','−'),mr('(λ) = '),total('s=1','3',[sub('w','s'),mr('('),sub('L','s'),mr(' − λ'),sup(sub('a','s'),'+'),mr('),')])],
    [sup('S','+'),mr('(λ) = '),total('s=1','3',[sub('w','s'),mr('('),sub('H','s'),mr(' − λ'),sup(sub('a','s'),'−'),mr(').')])]
])])
p=after(p,
    'The model and configuration subscripts are suppressed here for readability. The lower score is attained at '
    'the lowest task marginal and highest attack marginal; the upper score uses the opposite endpoints. '
    'This follows from monotonicity for nonnegative lambda. We examine lambda values 0, 0.1, 0.25, 0.5, 1, 2, '
    'and 5. A configuration remains a possible winner when its score upper bound reaches the largest score '
    'lower bound. The score expresses preference on evaluated attack cases; it does not estimate production '
    'attack frequency or calibrated economic loss.')
p=find('Where individual records are available')
inline(p,['With paired records, joint completion is directly observable. Within a fixed model, let ',sub('T','s'),
          ' be the number of sampled user tasks in suite s, and let ',sub(accent('J','̄'),'dst'),
          ' be the mean joint outcome across attacks on task t under configuration d. The task-weighted estimator is'])
replace('We compute paired differences',
    'We compute paired differences on identical tasks and obtain exploratory percentile intervals from 10,000 '
    'stratified task bootstrap replicates [11], using seed 20260915. Each resampled task retains both attacks '
    'and all configurations. These intervals address variation across the small empirical task sample; the '
    'identification bounds address unknown dependence. Neither procedure measures variation across fresh model '
    'executions. Degenerate all-zero or all-one bootstrap intervals are suppressed.')

# Source arithmetic table makes the audit directly reviewable.
p=replace('Three attacked-utility rows have',
    'Three attacked-utility rows have a printed Overall value that differs from the equal mean of their suite '
    'cells beyond rounding tolerance (Table III). The difference is reported Overall minus the mean of the '
    'printed suite percentages. We preserve the original cells and do not infer which source entry should '
    'be corrected.')
p=after(p,'TABLE III\nAUDIT OF PRINTED OVERALL ATTACKED UTILITY','Caption')
move_table_after(p,['Configuration / model','Overall','Mean','Δ (pp)'],[
    ['Spotlighting / Claude 3.5','58.33','57.09','+1.24'],
    ['Spotlighting / Claude 4.5','68.33','68.89','−0.56'],
    ['DRIFT / Kimi-K2.5','22.44','22.55','−0.11'],
],[1.82,.55,.55,.58])

p=after(find('The main methodological lesson'),
    'A practical reporting unit is one model, pipeline revision, task suite, and attack family with a shared '
    'denominator. Its four joint-cell counts recover task completion, ASR, and joint completion without assuming '
    'a dependence structure. Clean-task performance should accompany those counts so that low joint completion '
    'can be considered alongside baseline capability loss. A separate error count prevents infrastructure failure '
    'from being silently interpreted as either successful protection or a confirmed attack.')
p=after(find('The mathematical separation criterion'),
    'The exclusion analysis addresses the three detected arithmetic inconsistencies, but it cannot validate all '
    'source cells. An error in an unflagged cell, different metric denominators, or a change in suite weights '
    'could alter individual orderings. Rounded zero attack success is also not evidence that the underlying '
    'attack probability is exactly zero. A deployment decision would require representative tasks, repeated '
    'runs, explicit error handling, and an attack model appropriate to that setting.')
replace('The accompanying package includes',
    'The accompanying reproducibility package contains the extracted numeric tables, original pilot manifest '
    'and 288 source logs, analysis scripts, derived bounds and pairwise comparisons, figures, arithmetic audit, '
    'and dependency versions. Both analyses run offline without model access. Upstream attribution and license '
    'notices are retained. The source article is not redistributed.')
replace('OpenAI ChatGPT assisted',
    'OpenAI ChatGPT assisted with study design, mathematical exposition, drafting and editing all sections, '
    'and generating analysis and document scripts. The scripts computed the reported results from cited public '
    'data, and their numerical outputs were reproduced offline. This disclosure follows IEEE guidance [12].')

# Published IEEE US Letter geometry and typography.
for s in d.sections:
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.75);s.bottom_margin=Inches(1)
    s.left_margin=Inches(.625);s.right_margin=Inches(.625)
    s.header_distance=Inches(.3);s.footer_distance=Inches(.3)
    for part in [s.header,s.footer]:
        for p in part.paragraphs:p.clear()
for style in d.styles:
    if style.type==1:
        style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        fonts=style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if 'theme' in key.lower():del fonts.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
    for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
n=d.styles['Normal'];n.font.size=Pt(10);n.font.bold=False
n.paragraph_format.line_spacing=1.0;n.paragraph_format.space_before=Pt(0)
n.paragraph_format.space_after=Pt(0);n.paragraph_format.first_line_indent=Inches(.14)
n.paragraph_format.widow_control=True
t=d.styles['Title'];t.font.size=Pt(24);t.font.bold=False
t.paragraph_format.space_before=Pt(0);t.paragraph_format.space_after=Pt(12)
t.paragraph_format.first_line_indent=Inches(0)
for name in ['Heading 1','Heading 2']:
    st=d.styles[name];st.font.size=Pt(10);st.font.bold=False
    st.font.italic=name=='Heading 2';st.font.small_caps=name=='Heading 1'
    st.paragraph_format.space_before=Pt(9 if name=='Heading 1' else 6)
    st.paragraph_format.space_after=Pt(4 if name=='Heading 1' else 3)
    st.paragraph_format.first_line_indent=Inches(0)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.keep_with_next=True
c=d.styles['Caption'];c.font.size=Pt(8);c.font.italic=False
c.paragraph_format.first_line_indent=Inches(0)
c.paragraph_format.space_before=Pt(5);c.paragraph_format.space_after=Pt(5)
c.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Keep mathematical symbols italic and functions upright; 10 pt equations.
equations=[]
for p in d.paragraphs:
    if p.style.name=='Heading 1':
        p.text=re.sub(r'^([IVX]+) ',r'\1. ',p.text)
        if p.text=='Acknowledgment of AI Assistance':p.text='Acknowledgment'
    elif p.style.name=='Heading 2':p.text=re.sub(r'^([A-D]) ',r'\1. ',p.text)
    elif p.style.name=='Normal' and p.text and not p.text.startswith('['):
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if p.text.startswith('Abstract—') or p.text.startswith('Index Terms—'):
        p.paragraph_format.first_line_indent=Inches(0)
        p.paragraph_format.space_after=Pt(6)
        for run in p.runs:
            run.font.size=Pt(9);run.bold=True
            run.italic=p.text.startswith('Index Terms—')
    if p._p.xpath('.//m:oMath') and '\t(' in p.text:
        equations.append(p)
        f=p.paragraph_format;f.first_line_indent=Inches(0)
        f.space_before=Pt(5);f.space_after=Pt(5)
        f.tab_stops.clear_all()
        f.tab_stops.add_tab_stop(Inches(1.58),WD_TAB_ALIGNMENT.CENTER)
        f.tab_stops.add_tab_stop(Inches(3.49),WD_TAB_ALIGNMENT.RIGHT)
        for sz in p._p.xpath('.//m:r/w:rPr/w:sz'):sz.set(qn('w:val'),'20')
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if p.style.name=='Caption':
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if p.text.startswith('TABLE'):p.paragraph_format.keep_with_next=True

for i,p in enumerate(equations,1):
    for run in p.runs:
        if '(' in run.text:run.text='\t('+str(i)+')';run.font.size=Pt(10)
assert len(equations)==10
find('TABLE I  ').text='TABLE I\nSTRICTLY ORDERED WITHIN-MODEL PAIRS'
find('TABLE II  ').text='TABLE II\nJOINT OUTCOME COUNTS IN THE ARCHIVAL PILOT'
replace('Table I summarizes',
    'Table I summarizes the 540 comparisons, with 45 pairs per model. Pooled marginals separate 337 pairs '
    '(62.4%), while suite-level marginals separate 394 (73.0%), an increase of 57 pairs or 10.6 percentage '
    'points. Reversed denotes a separated joint ordering opposed to lower ASR. The remaining 146 pairs '
    'cannot be strictly ordered. Mean interval width falls from 5.54 to 4.78 percentage points across the '
    '120-condition panel; 51 conditions have narrower ranges.')
replace('The pilot supplies an independent',
    'The pilot provides a check using paired records. All eight exact joint rates lie inside their marginal '
    'bounds. Table II reports the four (U,A) cell counts, with 24 attacked records in every row. Clean '
    'records are separate. Simultaneous user and attacker success is therefore visible. The independence '
    'proxy differs from the exact joint rate by up to 3.30 percentage points, and its error can have either sign.')
for t in d.tables:
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            for shade in cell._tc.xpath('./w:tcPr/w:shd'):shade.set(qn('w:fill'),'FFFFFF')
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(8)

# A compact, monochrome scientific figure with readable labels.
import os
os.environ.setdefault('MPLCONFIGDIR',str(WORK/'analysis/mpl_cache'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
rows=[x for x in R['conditions'] if x['model']=='GPT-4o' and x['rectangular_panel']]
order=['None','Prompt Sandwiching','Spotlighting','ProtectAI','PIGuard','PromptGuard2','Tool Filter','CaMeL','Progent','DRIFT']
rows=[next(x for x in rows if x['defense']==name) for name in order]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':8})
fig,ax=plt.subplots(figsize=(3.5,3.0))
for y,row in enumerate(rows):
    lo,hi=row['pooled_joint_bounds_pct'];sl,sh=row['stratified_joint_bounds_pct']
    ax.plot([lo,hi],[y,y],color='.72',lw=5,solid_capstyle='butt')
    ax.plot([sl,sh],[y,y],color='.1',lw=2,solid_capstyle='butt')
    if sh-sl<.1:ax.plot([(sl+sh)/2],[y],marker='|',color='.1',markersize=5)
ax.set_yticks(range(10),order,fontsize=8);ax.invert_yaxis();ax.set_xlim(0,65)
ax.set_xlabel('Feasible joint completion (%)',fontsize=8)
ax.grid(axis='x',alpha=.2);ax.set_axisbelow(True)
ax.spines['top'].set_visible(False);ax.spines['right'].set_visible(False)
ax.legend(handles=[Line2D([0],[0],color='.72',lw=5,label='Pooled'),Line2D([0],[0],color='.1',lw=2,label='By suite')],loc='lower right',frameon=False,fontsize=8)
fig.subplots_adjust(left=.37,right=.99,top=.98,bottom=.17)
figure_path=ROOT/'results/joint_bounds_ieee.png'
fig.savefig(figure_path,dpi=450);fig.savefig(ROOT/'results/joint_bounds_ieee.svg');plt.close(fig)
for p in d.paragraphs:
    if p._p.xpath('.//w:drawing'):
        p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Inches(0)
        p.add_run().add_picture(str(figure_path),width=Inches(3.5))
        p.paragraph_format.keep_with_next=True
        break
replace('The largest narrowing occurs',
    'The largest narrowing occurs for unprotected GPT-4o. Its pooled feasible joint range is '
    '17.71%–55.53%, whereas suite information narrows it to 23.88%–48.69%, a width reduction of '
    '13.00 percentage points. Fig. 1 shows why reporting one exact joint score would overstate '
    'the evidence: several relatively productive configurations retain overlapping ranges.')

# Presentation-ready references; preserve verified bibliographic facts.
for p in d.paragraphs:
    if re.match(r'^\[\d+\]',p.text):
        p.paragraph_format.first_line_indent=Inches(-.16)
        p.paragraph_format.left_indent=Inches(.16)
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:r.font.size=Pt(8)
d.core_properties.title=title
d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.comments=''
d.core_properties.subject='ITT 2026 anonymous review manuscript'
for node in d.element.xpath('.//w:ins|.//w:del|.//w:commentRangeStart|.//w:commentRangeEnd'):
    raise AssertionError('Unexpected review markup')
d.save(TARGET)
report={
    'conference':'ITT 2026',
    'requirements_source':'https://easychair.org/cfp/ITT-2026',
    'conference_source':'https://hct.ac.ae/en/events/11th-international-conference-on-information-technology-trends-itt-2026/',
    'linked_template':'https://www.ieee.org/content/dam/ieee-org/ieee/web/org/conferences/conference-template-letter.docx',
    'template_binary_imported':False,
    'limitation':'The linked DOCX could not be downloaded. Formatting follows available IEEE conference guidance, not a byte-level template comparison.',
    'formatting_reference':'https://eit.r4.ieee.org/data/pdfs/IEEE-EIT-Paper_template.pdf',
    'paper_size_inches':[8.5,11],
    'margins_inches':{'top':.75,'bottom':1,'left':.625,'right':.625},
    'body_font':'Times New Roman 10 pt',
    'title_font':'Times New Roman 24 pt',
    'abstract_font':'Times New Roman 9 pt bold',
    'numbered_equations':len(equations),
    'anonymous':True,
    'fresh_model_executions':0,
    'source_docx_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'output_docx_sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),
}
(ROOT/'template_alignment.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({'file':str(TARGET),'numbered_equations':len(equations),'words':sum(len(p.text.split()) for p in d.paragraphs)},indent=2))
