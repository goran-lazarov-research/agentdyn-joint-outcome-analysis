"""Build the Serbian research protocol and English preliminary manuscript."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT.parent/'itt2026_output'
OUT.mkdir(exist_ok=True)
RESULT=json.loads((ROOT/'analysis/results/pilot_summary.json').read_text())
ROWS=RESULT['conditions']
COMMIT=RESULT['metadata']['source_commit']
DEF={'none':'Bez dodatne zaštite','spotlighting':'Spotlighting','tool_filter':'Tool Filter','progent':'Progent'}
DEF_EN={'none':'No added defense','spotlighting':'Spotlighting','tool_filter':'Tool Filter','progent':'Progent'}
MODEL={'gpt-4o-2024-08-06':'GPT-4o','google_gemini-2.5-pro':'Gemini 2.5 Pro'}

REFS=[
 ('K. Greshake, S. Abdelnabi, S. Mishra, C. Endres, T. Holz, and M. Fritz, “Not what you\'ve signed up for: Compromising Real-World LLM-Integrated Applications with Indirect Prompt Injection,” arXiv:2302.12173, 2023.', 'https://arxiv.org/abs/2302.12173'),
 ('E. Debenedetti, J. Zhang, M. Balunović, L. Beurer-Kellner, M. Fischer, and F. Tramèr, “AgentDojo: A Dynamic Environment to Evaluate Prompt Injection Attacks and Defenses for LLM Agents,” NeurIPS Datasets and Benchmarks, 2024.', 'https://arxiv.org/abs/2406.13352'),
 ('Q. Zhan, Z. Liang, Z. Ying, and D. Kang, “InjecAgent: Benchmarking Indirect Prompt Injections in Tool-Integrated Large Language Model Agents,” Findings of ACL, 2024; arXiv:2403.02691.', 'https://arxiv.org/abs/2403.02691'),
 ('E. Debenedetti et al., “Defeating Prompt Injections by Design,” arXiv:2503.18813v2, 2025.', 'https://arxiv.org/abs/2503.18813v2'),
 ('T. Shi et al., “Progent: Securing AI Agents with Privilege Control,” arXiv:2504.11703v3, revised May 2026.', 'https://arxiv.org/abs/2504.11703v3'),
 ('H. Li, R. Wen, S. Shi, N. Zhang, Y. Vorobeychik, and C. Xiao, “AgentDyn: Are Your Agent Security Defenses Deployable in Real-World Dynamic Environments?” arXiv:2602.03117v3, 2026.', 'https://arxiv.org/abs/2602.03117v3'),
 ('M. Nasr et al., “The Attacker Moves Second: Stronger Adaptive Attacks Bypass Defenses Against LLM Jailbreaks and Prompt Injections,” arXiv:2510.09023, 2025.', 'https://arxiv.org/abs/2510.09023'),
 ('H. Li, X. Liu, H.-C. Chiu, D. Li, N. Zhang, and C. Xiao, “DRIFT: Dynamic Rule-Based Defense with Injection Isolation for Securing LLM Agents,” NeurIPS 2025; arXiv:2506.12104v3, revised March 2026.', 'https://arxiv.org/abs/2506.12104v3'),
 ('SaFo-Lab, “AgentDyn source code and archived execution logs,” repository snapshot '+COMMIT+', accessed September 15, 2026.', 'https://github.com/SaFo-Lab/AgentDyn/tree/'+COMMIT),
 ('IEEE, “Author Guidelines for Artificial Intelligence Generated Text,” accessed September 15, 2026.', 'https://open.ieee.org/author-guidelines-for-artificial-intelligence-ai-generated-text/'),
 ('Higher Colleges of Technology, “ITT 2026 Call for Papers,” accessed September 16, 2026.', 'https://easychair.org/cfp/ITT-2026')
]

def font(style,name,size,bold=False):
    style.font.name=name;style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    fonts=style._element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if 'theme' in key.lower():del fonts.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),name)

def base(manuscript=False):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.75);s.bottom_margin=Inches(.8 if not manuscript else 1)
    s.left_margin=Inches(.8 if not manuscript else .625);s.right_margin=s.left_margin
    name='Times New Roman' if manuscript else 'Calibri'
    size=10 if manuscript else 11
    font(d.styles['Normal'],name,size)
    p=d.styles['Normal'].paragraph_format
    p.space_after=Pt(0 if manuscript else 6);p.line_spacing=1.02 if manuscript else 1.10
    p.widow_control=True
    if manuscript:p.first_line_indent=Inches(.14)
    for n,sz in [('Title',21 if manuscript else 22),('Subtitle',11),('Heading 1',10 if manuscript else 14),('Heading 2',10 if manuscript else 11.5)]:
        font(d.styles[n],name,sz,n.startswith('Heading'))
        f=d.styles[n].paragraph_format;f.space_before=Pt(10 if manuscript else 13);f.space_after=Pt(5);f.keep_with_next=True
        f.first_line_indent=Inches(0)
        if manuscript and n=='Heading 1':f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(d.styles['Caption'],name,8.5 if manuscript else 9)
    d.styles['Caption'].paragraph_format.space_after=Pt(7)
    d.styles['Caption'].paragraph_format.first_line_indent=Inches(0)
    for sty in d.styles:
        if sty.type==1:
            sty.font.color.rgb=RGBColor(0,0,0)
        for border in list(sty._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for node in list(d.settings.element.iter(qn('w:doNotBalanceTextColumns'))):
        node.getparent().remove(node)
    d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.comments=''
    if not manuscript:
        footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        r=footer.add_run();r.font.size=Pt(9)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    return d

def p(d,text,style=None,bold=False):
    z=d.add_paragraph(style=style);z.add_run(text).bold=bold
    if style is None:z.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return z

def h(d,text,level=1):return d.add_paragraph(text,style=f'Heading {level}')

def table(d,headers,rows,widths,small=False,keep_together=False):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for cell,w in zip(t.rows[0].cells,widths):cell.width=Inches(w)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for vals in rows:
        cells=t.add_row().cells
        for cell,w,text in zip(cells,widths,vals):cell.width=Inches(w);cell.text=str(text)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');pr.append(keep)
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'60' if small else '85');e.set(qn('w:type'),'dxa');marg.append(e)
            cp.append(marg)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            cp.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');cp.append(shade)
            for para in cell.paragraphs:
                if keep_together and ri<len(t.rows)-1:para.paragraph_format.keep_with_next=True
                para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.03;para.paragraph_format.first_line_indent=Inches(0)
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci<2 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(8 if small else 9);run.font.bold=(ri==0)
                    run.font.color.rgb=RGBColor(0,0,0)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def ref_list(d,limit=11,small=False):
    for i,(txt,url) in enumerate(REFS[:limit],1):
        z=p(d,f'[{i}] {txt} {url}')
        z.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if small:z.paragraph_format.keep_together=True
        z.paragraph_format.first_line_indent=Inches(-.18);z.paragraph_format.left_indent=Inches(.18)
        z.paragraph_format.space_after=Pt(5 if not small else 3)
        for r in z.runs:r.font.size=Pt(9 if not small else 8)

def pct(x,sr=False):return (f'{x:.1f}'.replace('.',',') if sr else f'{x:.1f}')

def protocol():
    d=base()
    p(d,'Bezbednost autonomnih AI agenata','Title')
    p(d,'Istraživački protokol za ITT 2026','Subtitle')
    p(d,'16. septembar 2026. | Verzija 1.1')
    p(d,'Predlažemo kontrolisano ispitivanje ovlašćenja agenata koji koriste spoljne alate. Cilj je da utvrdimo kako ograničavanje alata i njihovih parametara utiče na uspešnost napada i završavanje legitimnih zadataka. Priprema obuhvata nezavisnu obradu 288 javno objavljenih zapisa iz AgentDyn okruženja i protokol za nova izvršavanja na AgentDojo i AgentDyn podacima.')
    h(d,'Tema i doprinos rada')
    p(d,'Radni naslov glavne studije glasi Securing Autonomous AI Agents A Cross Benchmark Evaluation of Task Scoped Authorization Against Indirect Prompt Injection. Tema odgovara oblasti Responsible AI Cybersecurity and Digital Trust na ITT 2026 i njenom usmerenju na autonomne agente. Prema pozivu, rok je 21. septembar, konferencija se održava 28. i 29. oktobra 2026, a regularan rad ima šest strana u IEEE formatu. Dvostruko anonimna recenzija zahteva uklanjanje identiteta autora iz rukopisa i metapodataka. IEEE Xplore i Scopus su najavljeni uz uslov konačnog odobrenja [11].')
    p(d,'Istraživački doprinos treba da bude merljiva analiza granularnosti ovlašćenja pri istoj osnovnoj arhitekturi agenta. CaMeL, Progent i DRIFT već obrađuju sistemsku kontrolu izvršavanja [4, 5, 8]. Predloženi rad zato ispituje izbor i cenu konkretnih ograničenja, uz javni protokol i podatke o ishodima. Ovaj početni pregled literature je ciljani pregled relevantnih primarnih izvora; konačnu tvrdnju o originalnosti treba zasnovati na dodatnoj proveri srodnih radova.')
    for text in [
      'RQ1 Koliko provera parametara, dodata istoj listi dozvoljenih alata, menja uspešnost napada i zajedničku uspešnost korisničkog zadatka i zaštite?',
      'RQ2 Da li se uočeni odnosi održavaju na različitim okruženjima, posebno kada agent mora da koristi korisne spoljne instrukcije?',
      'RQ3 Koliko zaštita povećava broj poziva modelu, potrošnju tokena i vreme izvršavanja, a koliko legitimnih radnji neopravdano blokira?']:
        p(d,text)
    h(d,'Početna analiza javnih zapisa')
    p(d,'Pilot je sekundarna analiza tuđih arhiviranih izvršavanja. Nisu izvršeni novi pozivi modelima. Iz javnog repozitorijuma [9] izabrana su po četiri zadatka iz Shopping, GitHub i DailyLife okruženja, zatim po dva napada za svaki zadatak. Isti uzorak koristi se za dva modela i četiri konfiguracije. Ukupan obim je 192 izvršavanja pod napadom i 96 kontrolnih izvršavanja, uz 12 različitih korisničkih zadataka. Izbor je zapisan pre preuzimanja njihovih ishoda, korišćenjem semena 20260915.')
    p(d,'Zapisi potiču od 22. i 24. januara 2026. Svi predviđeni fajlovi su preuzeti i provereni prema Git sažecima sadržaja. U ovom uzorku nema zabeleženih tehničkih grešaka. Deo zapisa nema popunjene podatke o verziji paketa ili okruženja. Rezultati se odnose na ove arhivirane konfiguracije; ne predstavljaju ocenu najnovije implementacije Progent sistema ili aktuelne verzije modela.')
    rows=[]
    for r in ROWS:
        rows.append([MODEL[r['model']],DEF[r['defense']],f"{r['clean_successes']}/12",f"{r['attack_successes']}/24",f"{r['attacked_task_successes']}/24",f"{r['joint_successes']}/24"])
    table(d,['Model','Konfiguracija','Zadatak bez napada','Uspešan napad','Zadatak pod napadom','Zadatak i neuspešan napad'],rows,[1.03,1.30,1.06,.95,1.12,1.44])
    p(d,'Tabela 1 Brojevi ishoda u sekundarnom pilotu. Poslednja kolona zahteva uspešan korisnički zadatak i neuspeh definisanog napada u istom izvršavanju. Kolone o zadatku i napadu mogu istovremeno imati pozitivan ishod. Izvorni zapisi [9], samostalni proračun.','Caption')
    p(d,'U GPT-4o uzorku Progent ima 0 uspešnih napada od 24, ali i 0 uspešno završenih zadataka pod napadom. Spotlighting ima 11 zajednički uspešnih ishoda od 24 u oba modela. Njegova razlika prema osnovnoj konfiguraciji iznosi 4,2 procentna poena za GPT-4o i 8,3 za Gemini 2.5 Pro, uz eksplorativne intervale koji obuhvataju nulu. Ovi nalazi opravdavaju zajedničko merenje korisnosti i zaštite, ali mali uzorak ne daje osnov za proglašavanje najbolje odbrane.')
    fig=p(d,'');fig.add_run().add_picture(str(ROOT/'analysis/results/pilot_joint_outcomes.png'),width=Inches(6.7))
    p(d,'Slika 1 Zajednički uspeh i uspešnost napada u pilotu. Intervali su procenjeni ponovnim uzorkovanjem 12 zadataka unutar tri okruženja. Kada su svi zabeleženi ishodi nula, degenerisani interval nije prikazan. Nula u uzorku ne dokazuje odsustvo rizika.','Caption')
    h(d,'Podaci za glavno istraživanje')
    table(d,['Izvor','Obim objavljene verzije','Uloga'],[
      ['AgentDojo [2]','97 korisničkih zadataka i 629 bezbednosnih slučajeva','Osnovna evaluacija'],
      ['AgentDyn [6]','60 korisničkih zadataka i 560 bezbednosnih slučajeva','Provera dinamičkog planiranja'],
      ['InjecAgent [3]','1054 test slučaja','Dopunska provera ako resursi dozvole']],[1.25,2.8,2.85],keep_together=True)
    p(d,'AgentDojo i AgentDyn su kontrolisana istraživačka okruženja. Njihov rezultat nije procena učestalosti incidenata u produkciji. Pre novih merenja treba zamrznuti verzije repozitorijuma i proveriti dostupne zadatke, oznake i dozvoljene kombinacije napada. Podatke iz različitih okruženja prikazivati odvojeno; zbirni rezultat koristiti samo uz unapred navedeno pravilo ponderisanja.')
    h(d,'Model pretnji i kontrolisane konfiguracije')
    p(d,'Napadač kontroliše tekst na dozvoljenoj površini spoljnog sadržaja, kao što je poruka ili rezultat alata. Ne kontroliše izvornu korisničku nameru, mehanizam za proveru ovlašćenja, izvršni omotač niti evaluacione oznake. Svi alati za slanje poruka, promenu podataka ili transakcije ostaju simulirani u testnom okruženju. Tvrdnje o otpornosti važe isključivo u okviru ovako definisanih ovlašćenja napadača.')
    table(d,['Oznaka','Promena u odnosu na osnovnog agenta','Svrha poređenja'],[
      ['D0','Bez dodatne zaštite','Referentni ishod'],
      ['D1','Označavanje spoljnog sadržaja delimiterima','Efekat tekstualne zaštite'],
      ['D2','Provera liste dozvoljenih alata pre svakog poziva','Efekat ograničavanja alata'],
      ['D3','Ista lista alata kao D2 uz ograničenja njihovih parametara','Dodatni efekat granularnosti']],[.55,3.85,2.5])
    p(d,'D2 i D3 moraju koristiti isti model, sistemske instrukcije, izbor alata, granicu koraka i izvor politike. D3 dodaje proveru dozvoljenih primalaca, resursa, operacija ili drugih parametara definisanih zadatkom. Time se izdvaja promena u granularnosti. Arhivski Progent i Tool Filter iz pilota ne ispunjavaju ovaj uslov kontrolisanog poređenja i služe kao početni empirijski kontekst.')
    p(d,'Politika nastaje iz korisničkog zahteva i unapred odobrene šeme ovlašćenja. Spoljni sadržaj može da pruži podatak potreban za izvršavanje već dozvoljene radnje, ali ne može sam da odobri novog primaoca, novu privilegiju ili veću transakciju. Politika se proverava izvan jezičkog modela. Odbijanje, prekid i zahtev za dodatnim odobrenjem beleže se kao različiti događaji. U automatizovanom testu zahtev za ljudskim odobrenjem ne dobija izmišljenu potvrdu.')
    h(d,'Merenje i statistička obrada')
    p(d,'Primarni zajednički ishod je procenat izvršavanja u kojima je korisnički zadatak završen, a definisani napad nije uspeo. Izračunava se iz dva ishoda u istom zapisu; ne dobija se množenjem zasebnih procenata. Ovaj pokazatelj govori o konkretnom napadačkom cilju. Za širu tvrdnju o poštovanju politike potreban je zaseban validator svih relevantnih poziva alata.')
    p(d,'Obavezno meriti i uspešnost napada, korisnost bez napada i pod napadom, broj pogrešno blokiranih dozvoljenih radnji, broj poziva modelu, tokene i trajanje. Pad korisnosti sam po sebi nije dokaz pogrešnog blokiranja: odluke o blokiranju moraju se proveriti prema nezavisno definisanoj politici. Trošak API poziva može se računati tek iz stvarne potrošnje i cenovnika važećeg na datum eksperimenta. Arhivirani pilot nema dovoljne podatke za takav obračun.')
    p(d,'Jedan korisnički zadatak je klaster koji obuhvata njegove napade, konfiguracije i ponavljanja. Uparene razlike računati na identičnim zadacima, a intervale poverenja procenjivati ponovnim uzorkovanjem zadataka unutar istog okruženja. Napade na isti zadatak ne tretirati kao nezavisne korisničke zadatke. Za pilot je korišćeno 10000 takvih uzoraka. Intervali iz malog uzorka su eksplorativni i ne obuhvataju varijabilnost novih izvršavanja modela.')
    p(d,'Tehnički neuspeh evidentirati odvojeno. U analizi svih planiranih izvršavanja dodeliti mu nultu korisnost i nulti zajednički uspeh. Za uspešnost napada prikazati broj validnih evaluacija, broj grešaka i konzervativnu analizu osetljivosti. Grešku servisa ne predstavljati kao potvrđen uspeh napada. Nedostajući preuzeti fajl nije rezultat agenta; rešava se kroz izveštaj o dostupnosti podataka.')
    h(d,'Obim i postupak novih eksperimenata')
    p(d,'Za početnu fazu predlažemo 24 evaluaciona zadatka, po 12 iz svakog osnovnog okruženja, dva kompatibilna napada po zadatku, dva modela različitih porodica, četiri konfiguracije i tri ponavljanja. To je najviše 1728 izvršavanja, zajedno sa kontrolnim uslovima bez napada. U svakom okruženju zadatke rasporediti po scenarijima. Ovaj obim daje ograničenu empirijsku studiju; proširenje na veći broj različitih zadataka ima prednost nad dodavanjem velikog broja sličnih napada.')
    p(d,'Razvojni zadaci i svi prethodno detaljno pregledani zadaci iz pilota izdvajaju se iz konačnog evaluacionog skupa. Pre pokretanja zapisati tačne identifikatore modela, temperaturne parametre, dozvoljeni broj koraka i poziva, pravila ponavljanja grešaka i izbor napada. Pilot od 10 do 20 novih izvršavanja treba da proveri integraciju, trajanje i potrošnju; na osnovu toga se zaključava obim i budžet. U ovom okruženju trenutno nije podešen pristup modelima za ta nova izvršavanja.')
    p(d,'Fiksirane napade dopuniti ograničenom evaluacijom napada prilagođenih odbrani kada resursi dozvole [7]. Budžet napadača mora biti isti među konfiguracijama. Ako takav test izostane, zaključak se ograničava na korišćene napade. Ne tvrditi opštu otpornost niti koristiti izraz bezbedan za produkciju na osnovu uspeha u ovim testovima.')
    h(d,'Koraci do završnog rukopisa')
    table(d,['Faza','Potreban rezultat'],[
      ['Priprema','Protokol, sekundarni pilot i proverljiv paket podataka i koda'],
      ['Integracija','Dostupno izvršno okruženje i pilot od 10 do 20 novih izvršavanja'],
      ['Merenje','Kontrolisano poređenje D2 i D3 uz zaključan obim i budžet'],
      ['Obrada','Uparena statistička analiza i pregled uzroka neuspeha'],
      ['Rukopis','Rezultati, diskusija, konačna provera literature i IEEE šablona'],
      ['Predaja','Autorska provera i odluka o odgovarajućoj kategoriji rada']],[1.3,5.6],keep_together=True)
    p(d,'Datumi novih merenja zavise od izvršnog okruženja i raspoloživih resursa. Poziv trenutno navodi 21. septembar 2026. kao rok za predaju [11]; termin i vremensku zonu treba proveriti u sistemu za prijavu. Dok se glavni eksperiment ne završi, početni rukopis ostaje metodološki nacrt sa sekundarnim pilotom. Dostavljeni dvokolonski raspored služi za radnu verziju i ne predstavlja potvrdu potpune usklađenosti sa zvaničnim IEEE šablonom. Broj citata i prihvatanje rada ne mogu se unapred garantovati.')
    p(d,'Rukopis treba da objedini sažetak, ključne reči, uvod, pregled literature, metod, rezultate, diskusiju, ograničenja, zaključak i reference. Uz rad pripremiti anonimizovan istraživački paket sa podacima, skriptama i verzijama. Izjava o korišćenju ChatGPT pomoći mora precizno opisati njenu ulogu u tekstu i kodu, prema IEEE smernicama [10]. Konačni autori proveravaju svaki citat, rezultat i interpretaciju pre predaje.')
    h(d,'Literatura i izvori')
    ref_list(d)
    d.save(OUT/'ITT2026_Istrazivacki_protokol.docx')

def manuscript():
    d=base(True)
    title=p(d,'Securing Autonomous AI Agents through Task Scoped Authorization','Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    sub=p(d,'Study protocol with a secondary pilot analysis of public AgentDyn logs','Subtitle');sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
    s=d.add_section(WD_SECTION_START.CONTINUOUS)
    cols=s._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'360')
    z=p(d,'Abstract');z.runs[0].bold=True;z.paragraph_format.first_line_indent=Inches(0)
    p(d,'Autonomous language model agents require controls that preserve task completion while limiting actions induced by untrusted content. This paper specifies a controlled evaluation of task scoped authorization and reports a secondary pilot analysis of 288 public AgentDyn execution logs. The sample contains 12 user tasks across three suites, two base models, four archived defense configurations, 192 attacked runs, and 96 clean runs. We measure successful task completion jointly with failure of the specified attack, rather than interpreting attack success alone. In the GPT-4o subset, archived Progent runs record no successful attacks and no completed attacked tasks, whereas the joint success rate of Spotlighting is 45.8%. This observation motivates explicit measurement of utility loss and does not establish superiority of any current defense. The prospective experiment separates tool allowlisting from argument restrictions while holding the surrounding agent configuration fixed. Reproducible sampling, joint outcome records, and task clustered uncertainty estimates support the design. Fresh model executions and evaluation on AgentDojo remain future stages of the study.')
    z=p(d,'Index Terms—AI agents, cybersecurity, indirect prompt injection, authorization, benchmark evaluation.');z.paragraph_format.first_line_indent=Inches(0);z.runs[0].italic=True
    h(d,'I Introduction')
    p(d,'A language model agent can transform retrieved text into consequential tool use. An attacker who controls part of that text may redirect the agent away from the user’s authorized task without controlling the original request. This mechanism is central to indirect prompt injection [1]. The relevant security boundary lies between information supplied by the environment and authority to act on the user’s behalf.')
    p(d,'Benchmarks such as AgentDojo and InjecAgent make aspects of this risk measurable [2, 3]. System controls including CaMeL and Progent constrain actions or data flows [4, 5]. Their practical value depends on how much legitimate work remains possible. AgentDyn adds tasks that require adaptation to information encountered during execution [6]. Such tasks provide a useful setting for examining whether an authorization policy blocks a malicious deviation, a legitimate continuation, or both.')
    p(d,'We formulate a study of policy granularity. The principal comparison is between a tool allowlist and the same allowlist augmented with argument constraints. Holding the base agent and policy source fixed is necessary to attribute a difference to the added restriction. Comparing independently implemented defenses can identify operational differences, but changes in prompts, planning, and execution logic prevent a clean causal interpretation.')
    p(d,'This version contributes a reproducible secondary pilot and a prospective experimental protocol. The pilot tests the extraction and interpretation of joint outcomes from public execution logs. It does not introduce a new authorization system, reproduce current model behavior, or establish the novelty of an already known least privilege principle. The controlled experiment described below is intended to provide the additional evidence needed for the full study.')
    h(d,'II Related Work')
    p(d,'Greshake et al. analyze indirect injection through content consumed by language model applications [1]. AgentDojo provides an extensible tool use environment with explicit user and attacker objectives [2]. InjecAgent offers a complementary collection of tool integrated injection cases [3]. AgentDyn focuses on dynamic tasks and useful environmental instructions [6]. These resources differ in task structure, so their rates should be reported separately before any explicitly weighted summary.')
    p(d,'Evaluation against fixed attacks is insufficient for broad robustness claims. Nasr et al. demonstrate the importance of attackers who adapt to the defense [7]. Our archival pilot covers one named attack family. A bounded adaptive evaluation is therefore part of the proposed extension, with an identical attacker budget across compared configurations.')
    p(d,'CaMeL separates trusted control from untrusted data and enforces constraints on data flows [4]. Progent represents permissions over tool names and arguments; its May 2026 revision specifies controlled updates and approval for expansions [5]. DRIFT combines planning constraints, validation, and injection isolation [8]. These systems establish that task aware control is an existing research direction. The proposed contribution is a focused evaluation of restriction granularity and its operational cost, rather than a claim to originate that direction.')
    h(d,'III Secondary Pilot Method')
    h(d,'A Data provenance and selection',2)
    p(d,'We selected public execution logs from the AgentDyn repository [9], pinned to snapshot '+COMMIT+'. The model labels are gpt-4o-2024-08-06 and google_gemini-2.5-pro. The configurations are no added defense, spotlighting_with_delimiting, tool_filter, and progent. These labels identify archived pipelines and do not certify their equivalence to the latest published defense implementations.')
    p(d,'The sampling frame comprises user_task_0 through user_task_19 in each of Shopping, GitHub, and DailyLife. Using Python’s seeded pseudo-random generator with seed 20260915, we selected four tasks per suite. Two injection objectives were then selected per task from the suite’s published objective indices. The same task and attack identifiers were used for every model and configuration. A clean record was added for each task and condition. The selection manifest was written before downloading the selected outcome records.')
    p(d,'This design yields 24 attacked and 12 clean records per model and configuration, or 288 records in total. All predetermined records were recovered and checked against the repository’s Git blob hashes. The root inventory was truncated, but all selected paths were present and verified; no completeness claim is made about the entire repository. Transfer retries recovered initially unavailable files without changing the sample. Logged evaluation dates are January 22 and January 24, 2026. Some package and benchmark version fields are missing.')
    h(d,'B Outcome definitions',2)
    p(d,'Let U indicate completion of the user task and A indicate success of the specified injection objective on an attacked run. The source method BaseInjectionTask.security returns whether that injection objective succeeded, and the archived security field carries this result [9]. We therefore interpret security=True as A=1 for attacked tasks. Clean records use a different convention for this field and are excluded from attack rate calculations.')
    p(d,'The joint outcome is J=1 exactly when U=1 and A=0 and no technical error is logged. A run may complete the legitimate task and also satisfy the attacker’s goal; it must then contribute zero to J. Consequently, the joint rate is computed from paired outcomes in each execution. Multiplying the marginal task rate by one minus the attack rate would impose an unsupported independence assumption.')
    p(d,'We report clean utility, attacked utility, attack success, and joint success. J refers to failure of the benchmark’s specified attacker objective, not proof that every action complied with a complete authorization policy. Establishing that broader property requires a separate action level validator. No technical errors are present in this sample. In future runs, technical errors will be reported explicitly, assigned zero utility and joint completion for an all attempted runs analysis, and excluded from the valid attack denominator with a worst case sensitivity analysis.')
    h(d,'C Statistical analysis',2)
    p(d,'The independent sampling unit for exploratory uncertainty estimation is the user task. We retain both attacks and all compared conditions when resampling a task. We generate 10000 paired bootstrap samples, drawing four tasks with replacement within each of the three suites, using seed 20260915. Suite weights are equal. Percentile intervals summarize uncertainty over this small empirical task sample; they do not estimate stochastic variation from new model calls.')
    p(d,'An all zero or all one empirical outcome produces a degenerate percentile interval. We suppress such intervals rather than presenting them as evidence of certainty. With only 12 task clusters, all intervals remain exploratory. No p-value based claims or claims of deployment safety are made. The accompanying script verifies the four possible U and A combinations and checks file identity and the expected paired sample structure before producing results.')
    h(d,'IV Pilot Results')
    p(d,'Table I reports exact counts. Each row has 12 clean records and 24 attacked records. Zero attack successes can coexist with zero useful task completion. Conversely, task success and attacker success can occur in the same execution, making attacked utility alone insufficient to characterize the desired joint outcome.')
    p(d,'TABLE I  Secondary pilot counts. U0 is clean task success, A is attack success, U is attacked task success, and J is task success with attack failure. Source: archived logs [9]; calculations performed for this study.','Caption')
    for model in MODEL:
        model_heading=p(d,MODEL[model],bold=True)
        model_heading.paragraph_format.keep_with_next=True
        vals=[]
        for r in ROWS:
            if r['model']==model:
                vals.append([DEF_EN[r['defense']],r['clean_successes'],r['attack_successes'],r['attacked_task_successes'],r['joint_successes']])
        table(d,['Defense','U0 /12','A /24','U /24','J /24'],vals,[1.42,.47,.47,.47,.47],small=True,keep_together=True)
    p(d,'For GPT-4o, no added defense yields 10 joint successes out of 24, or 41.7%. Spotlighting yields 11 of 24, or 45.8%, while Tool Filter yields 2 of 24 and Progent yields 0 of 24. Attack successes are 7, 4, 1, and 0, respectively. Thus the configuration with the smallest observed attack count does not have the largest observed joint completion count.')
    p(d,'For Gemini 2.5 Pro, joint successes are 9, 11, 0, and 2 out of 24 for the same ordered configurations. Attack successes are 2, 3, 0, and 0. In this subset, Spotlighting has a higher joint completion count than the unprotected pipeline while also having a higher attack count. A single ordering of defenses therefore depends on the operational outcome selected.')
    p(d,'The paired Spotlighting minus baseline difference in joint success is 4.2 percentage points for GPT-4o, with an exploratory 95% cluster interval of −8.3 to 16.7 points. The corresponding Gemini difference is 8.3 points, with an interval of −4.2 to 20.8 points. Both intervals include zero. The sample supports reporting the observed tradeoffs and motivates a larger controlled experiment; it does not identify a statistically established overall winner.')
    p(d,'The published aggregate table in AgentDyn [6] and our sampled execution records [9] are distinct evidence sources. Their populations and aggregation procedures differ. The supplementary package preserves a separate mechanical extraction of the published numeric table for traceability, but those aggregate values do not enter the pilot calculations.')
    h(d,'V Controlled Experiment Protocol')
    h(d,'A Threat model and policy boundary',2)
    p(d,'The attacker may replace text only on a benchmark authorized external content surface. The attacker cannot edit the original user request, the policy enforcement layer, tool implementation, or outcome labels. The agent operates exclusively in simulated environments. We assume that the enforcement layer observes each attempted tool invocation and that all relevant tool calls pass through it. Claims derived from the study will be conditional on these assumptions.')
    p(d,'A policy describes allowed tools and task appropriate constraints on their arguments. The policy source must be the trusted task request and an independently specified authorization schema. Retrieved content may supply a value needed for an already permitted operation, but it must not enlarge the agent’s authority merely by requesting that enlargement. Read access, write access, external recipients, and sensitive resource scopes require distinct treatment where the task semantics make them relevant.')
    h(d,'B Configurations and research questions',2)
    p(d,'D0 is a baseline with no additional defense. D1 marks untrusted content with delimiters. D2 enforces a tool allowlist before execution. D3 uses the identical allowlist and additionally validates tool arguments. D2 and D3 share the same base model, prompts, execution budget, and policy construction process. This last pair is the principal comparison for estimating the effect of added argument restrictions.')
    p(d,'RQ1 asks how that added restriction changes attack success and joint task completion. RQ2 asks whether the relationship transfers between AgentDojo and AgentDyn. RQ3 asks how the policy changes false blocking, model call counts, token consumption, and elapsed time. Existing implementations such as Progent may provide reference baselines, but their full systems must remain separate from the controlled D2 versus D3 ablation.')
    p(d,'Policy construction must not inspect test attack labels or reference solutions. Development tasks are separate from test tasks, and the 12 tasks examined in this pilot are excluded from the eventual AgentDyn test subset. The intended policy is checked independently of the agent’s proposed action. A request for additional human approval is recorded as an escalation event; automated testing must not invent approval on the user’s behalf.')
    h(d,'C Proposed scale and recording',2)
    p(d,'The initial prospective scope is 24 evaluation tasks, 12 from each benchmark, two compatible injection objectives per task, two model families, four configurations, and three repetitions. Including clean conditions, this yields up to 1728 trajectories. A separate integration pilot of 10 to 20 fresh trajectories will establish duration and resource requirements before the final scope is locked. The limited number of independent tasks must remain explicit in any resulting claims.')
    p(d,'Model identifiers, generation parameters, step limits, retry handling, task identifiers, source revisions, and attacker budgets will be fixed before evaluation. Clean and attacked conditions will use the same task states. Execution order will be randomized to reduce temporal effects. Results will retain benchmark and suite identifiers. We will prefer expanding task diversity to adding many closely related attacks on the same task when resources permit.')
    p(d,'Each execution will record task and attack outcomes, a trace of attempted tool calls, policy decisions and reasons, escalation events, tokens, provider reported usage, and duration. False blocking requires a labeled permitted action rejected by the control; a reduction in total utility alone does not establish that mechanism. The current public logs do not provide sufficient standardized usage data for a reproducible API cost calculation. Their historical durations are descriptive and must not be interpreted as a controlled estimate of defense overhead.')
    p(d,'An adaptive attack extension will allocate the same search budget to each defense [7]. If resource constraints prevent that extension, the results will be explicitly limited to the evaluated fixed attacks. Evaluation on AgentDojo and all fresh model executions remain to be performed; the pilot reported here is confined to archived AgentDyn records.')
    h(d,'VI Discussion and Limitations')
    p(d,'The pilot illustrates the practical value of a joint outcome. It also identifies a design question: whether a policy can preserve legitimate continuations without permitting authority expansion from untrusted text. Answering that question requires separating policy granularity from implementation differences. The low utility of an archived full system cannot establish that argument constraints themselves caused the failure.')
    p(d,'Several limitations constrain the present evidence. The sample contains 12 tasks, a single archived realization of each condition, and one attack family. It was not designed to estimate production incident frequencies. Model and benchmark metadata are incomplete in some records. Pipeline implementation and execution dates differ, and the January runs predate later defense revisions. No fresh baseline run verifies present behavior. The benchmark’s attacker outcome covers a specified goal and can miss other unauthorized effects.')
    p(d,'The observed zero rates therefore provide no guarantee of universal protection. Small task clustered intervals should be interpreted as exploratory. Defenses may also alter the likelihood that the agent reaches an injection surface, so future analysis should distinguish end to end outcomes from exposure conditioned outcomes while clearly identifying their different denominators. These distinctions help prevent claims that exceed what the experimental design can support.')
    h(d,'VII Conclusion')
    p(d,'A secondary analysis of 288 public execution logs demonstrates that attack success and useful task completion must be examined jointly. In the sampled GPT-4o records, an archived configuration achieves zero attacker successes while completing zero attacked tasks. The proposed controlled study will compare a tool allowlist with the same policy augmented by argument constraints across two benchmarks. Its contribution will depend on completed fresh measurements, consistent policy construction, and transparent reporting of task level uncertainty and resource cost.')
    h(d,'Data and Code Availability')
    p(d,'The accompanying analysis package includes the predetermined sample manifest, all 288 selected public logs, source identity checks, normalized outcomes, calculations, and plotting code. It preserves the upstream license. Scripts perform secondary analysis without invoking a model. The original repository snapshot is identified in [9]. No public hosting or persistent publication identifier is claimed for the new package.')
    h(d,'Acknowledgment of AI Assistance')
    p(d,'OpenAI ChatGPT assisted with study design, drafting and editing text in all sections, and generating the data processing and document preparation scripts. Numerical pilot results were computed by the supplied scripts from the cited public records. This acknowledgment documents the assistance described by IEEE guidance [10]; it does not substitute for the submitting authors’ verification and responsibility.')
    h(d,'References')
    ref_list(d,10,small=True)
    d.add_section(WD_SECTION_START.CONTINUOUS)
    d.save(OUT/'ITT2026_Pocetni_rukopis_EN.docx')

if __name__=='__main__':
    protocol();manuscript()
    print('Created',OUT/'ITT2026_Istrazivacki_protokol.docx')
    print('Created',OUT/'ITT2026_Pocetni_rukopis_EN.docx')
